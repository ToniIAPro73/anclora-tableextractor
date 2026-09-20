"""Export layer: xlsx / csv / json with origin-page + confidence metadata.
Supports single-document and batch (multi-document) exports."""
import io
import json
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _grid(table):
    ncols = table["num_columnas"]
    nrows = table["num_filas"]
    values = [["" for _ in range(ncols)] for _ in range(nrows)]
    meta = [[{} for _ in range(ncols)] for _ in range(nrows)]
    for cell in table["cells"]:
        r, c = cell["fila"], cell["columna"]
        if r < nrows and c < ncols:
            values[r][c] = cell.get("valor", "")
            meta[r][c] = {"score": cell.get("score_confianza", 1.0), "page": cell.get("pagina", 1)}
    return table["columnas"], values, meta


def export_json(document, tables):
    payload = {
        "documento": {
            "id": document["id"],
            "nombre_archivo": document["nombre_archivo"],
            "num_paginas": document["num_paginas"],
            "estado": document["estado"],
        },
        "tablas": [],
    }
    for t in tables:
        cols, values, meta = _grid(t)
        rows = []
        for r in range(len(values)):
            row = {}
            for c in range(len(cols)):
                row[cols[c]] = {
                    "valor": values[r][c],
                    "score_confianza": meta[r][c].get("score", 1.0),
                    "pagina_origen": meta[r][c].get("page", 1),
                }
            rows.append(row)
        payload["tablas"].append({
            "id": t["id"],
            "pagina_origen": t["pagina_origen"],
            "source_pages": t.get("source_pages", [t["pagina_origen"]]),
            "extraction_method": t["extraction_method"],
            "columnas": cols,
            "filas": rows,
        })
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _csv_escape(v):
    v = "" if v is None else str(v)
    if any(ch in v for ch in [",", '"', "\n"]):
        return '"' + v.replace('"', '""') + '"'
    return v


def export_csv(document, tables):
    buf = io.StringIO()
    for ti, t in enumerate(tables):
        cols, values, meta = _grid(t)
        buf.write(f"# Tabla {ti + 1} | paginas: {t.get('source_pages', [t['pagina_origen']])} | metodo: {t['extraction_method']}\n")
        header = []
        for c in cols:
            header += [c, f"{c}__score", f"{c}__pagina"]
        buf.write(",".join(_csv_escape(h) for h in header) + "\n")
        for r in range(len(values)):
            line = []
            for c in range(len(cols)):
                line += [values[r][c], str(meta[r][c].get("score", "")), str(meta[r][c].get("page", ""))]
            buf.write(",".join(_csv_escape(x) for x in line) + "\n")
        buf.write("\n")
    return buf.getvalue().encode("utf-8")


_GREEN = PatternFill("solid", fgColor="D1FAE5")
_AMBER = PatternFill("solid", fgColor="FEF3C7")
_RED = PatternFill("solid", fgColor="FEE2E2")
_HEADER_FILL = PatternFill("solid", fgColor="0EA5E9")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _add_doc_sheets(wb, document, tables, prefix=""):
    base = (document["nombre_archivo"] or "doc").rsplit(".", 1)[0]
    for ti, t in enumerate(tables):
        cols, values, meta = _grid(t)
        title = f"{prefix}{base}_{ti + 1}"[:31]
        # openpyxl forbids some chars in sheet titles
        for ch in "[]:*?/\\":
            title = title.replace(ch, "-")
        ws = wb.create_sheet(title=title or f"Tabla_{ti + 1}")
        for c, name in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        pcol = len(cols) + 1
        h = ws.cell(row=1, column=pcol, value="_pagina_origen")
        h.font = _HEADER_FONT
        h.fill = _HEADER_FILL
        for r in range(len(values)):
            for c in range(len(cols)):
                cell = ws.cell(row=r + 2, column=c + 1, value=values[r][c])
                score = meta[r][c].get("score", 1.0)
                cell.fill = _GREEN if score >= 0.9 else _AMBER if score >= 0.6 else _RED
            page_vals = sorted({meta[r][c].get("page", 1) for c in range(len(cols))})
            ws.cell(row=r + 2, column=pcol, value=", ".join(str(p) for p in page_vals))


def export_xlsx(document, tables):
    wb = Workbook()
    wb.remove(wb.active)
    _add_doc_sheets(wb, document, tables)
    if not wb.sheetnames:
        wb.create_sheet("Vacio")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_batch_xlsx(items):
    """items: list of (document, tables). One workbook, sheets per doc/table."""
    wb = Workbook()
    wb.remove(wb.active)
    for i, (document, tables) in enumerate(items, start=1):
        _add_doc_sheets(wb, document, tables, prefix=f"{i}_")
    if not wb.sheetnames:
        wb.create_sheet("Vacio")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_batch_zip(items, fmt):
    """items: list of (document, tables). ZIP with one file per document."""
    buf = io.BytesIO()
    seen = {}
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for document, tables in items:
            base = (document["nombre_archivo"] or "doc").rsplit(".", 1)[0]
            name = f"{base}.{fmt}"
            seen[name] = seen.get(name, 0) + 1
            if seen[name] > 1:
                name = f"{base}_{seen[name]}.{fmt}"
            data = export_csv(document, tables) if fmt == "csv" else export_json(document, tables)
            zf.writestr(name, data)
    return buf.getvalue()
